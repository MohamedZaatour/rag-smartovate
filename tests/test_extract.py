# ============================================================
#  Tests automatiques de la chaine d'extraction (US 1.1)
#  Lancer depuis 02-Developpement avec :  pytest
# ============================================================

import docx
import pymupdf
import pytest

from ingestion.extract import nettoyer_texte, traiter_document

# --- nettoyer_texte : fonction pure, facile a tester ---

def test_nettoyer_reduit_espaces_multiples():
    assert nettoyer_texte("Bonjour    le   monde") == "Bonjour le monde"


def test_nettoyer_reduit_sauts_de_ligne_multiples():
    assert nettoyer_texte("a\n\n\n\n\nb") == "a\n\nb"


# --- DOCX : les tableaux doivent etre lus (correction #1) ---

def test_docx_lit_les_tableaux(tmp_path):
    chemin = tmp_path / "doc.docx"
    d = docx.Document()
    d.add_paragraph("Intro")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Critere"
    t.rows[0].cells[1].text = "Etat"
    d.save(str(chemin))

    res = traiter_document(str(chemin))
    assert res["type"] == "docx"
    assert res["nb_pages"] == 1
    assert "Critere | Etat" in res["pages"][0]["texte"]


# --- PDF : structure page par page (correction #2) ---

def test_pdf_page_par_page(tmp_path):
    chemin = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Page une")
    doc.new_page().insert_text((72, 72), "Page deux")
    doc.save(str(chemin))
    doc.close()

    res = traiter_document(str(chemin))
    assert res["nb_pages"] == 2
    assert res["pages"][0]["page"] == 1
    assert res["pages"][1]["page"] == 2


# --- Alerte document vide (correction #3) ---

def test_document_vide_leve_une_erreur(tmp_path):
    chemin = tmp_path / "vide.pdf"
    doc = pymupdf.open()
    doc.new_page()  # page blanche, aucun texte
    doc.save(str(chemin))
    doc.close()

    with pytest.raises(ValueError):
        traiter_document(str(chemin))


# --- Format non supporte ---

def test_format_non_supporte_leve_une_erreur(tmp_path):
    chemin = tmp_path / "fichier.txt"
    chemin.write_text("bonjour")

    with pytest.raises(ValueError):
        traiter_document(str(chemin))
